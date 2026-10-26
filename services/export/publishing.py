"""Professional A4 PDF and editable DOCX export, with strict preflight.

Integration note: call only after the authenticated API validates tenant ownership,
review approval, and teacher-only audience authorization. Never expose these
functions as a public unauthenticated download handler.
"""
from __future__ import annotations

from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any
import re
from xml.etree import ElementTree as ET

from services.rendering.document import DocumentValidationError, validate_document
from services.rendering.omml import latex_to_omml


def _approved(document: dict[str, Any]) -> None:
    if document.get("review_state") not in {"approved", "published"}:
        raise DocumentValidationError("An educator must approve the version before publication")


def _branding(branding: dict[str, str] | None) -> dict[str, str]:
    brand = branding or {}
    return {"school_name": str(brand.get("school_name") or "SomaLoom School"),
            "footer": str(brand.get("footer") or "Prepared with SomaLoom")}


def export_docx(
    document: dict[str, Any], branding: dict[str, str] | None = None,
    *, audience: str = "student", require_approval: bool = True
) -> bytes:
    """Return editable Office DOCX; equations are native Office Math, not PNGs."""
    if require_approval:
        _approved(document)
    data = validate_document(document, audience=audience)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    b = _branding(branding)
    out = Document()
    sec = out.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
    normal = out.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = out.styles[style_name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(19, 53, 65)
    head = sec.header.paragraphs[0]
    head.text = b["school_name"]
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot = sec.footer.paragraphs[0]
    foot.text = b["footer"] + "  •  Page "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    foot._p.append(field)
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    out.add_paragraph(b["school_name"], style="Subtitle")
    out.add_heading(data["title"], 0)
    out.add_paragraph("TEACHER EDITION" if audience == "teacher" else "LEARNER EDITION")
    for section in data["body"]["sections"]:
        if section.get("title"):
            out.add_heading(str(section["title"]), level=1)
        for block in section["blocks"]:
            kind = block["type"]
            if kind == "page_break":
                out.add_page_break()
            elif kind == "heading":
                out.add_heading(str(block.get("text", "")), level=2)
            elif kind == "list":
                for line in str(block.get("text", "")).splitlines():
                    if line.strip():
                        out.add_paragraph(line.strip(), style="List Bullet")
            elif kind == "table":
                rows = block["rows"]
                table = out.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        table.cell(i, j).text = cell
            elif kind == "equation":
                p = out.add_paragraph()
                math_xml = ET.tostring(latex_to_omml(block["latex"]), encoding="unicode")
                p._p.append(parse_xml(math_xml))
            elif kind == "reference":
                out.add_paragraph(str(block.get("text", "")), style="Quote")
            elif kind in {"paragraph", "question", "activity", "answer", "marking_guide"}:
                out.add_paragraph(str(block.get("text", "")))
            else:
                raise DocumentValidationError("Unrenderable block: " + kind)
            if block.get("citations"):
                out.add_paragraph("Evidence IDs: " + ", ".join(block["citations"]), style="Caption")
    stream = BytesIO()
    out.save(stream)
    return stream.getvalue()


def _matrix_rows(latex: str) -> list[list[str]] | None:
    expression = re.fullmatch(
        r"\\begin\{(matrix|pmatrix|bmatrix)\}(.*?)\\end\{\1\}", latex.strip(), re.S
    )
    if not expression:
        return None
    rows = [row.split("&") for row in re.split(r"\\\\", expression.group(2))]
    if len({len(row) for row in rows}) != 1:
        raise DocumentValidationError("Matrix has uneven columns")
    return [[cell.strip() for cell in row] for row in rows]


def _svg_equation(latex: str, max_width: float):
    """Matplotlib mathtext creates vector SVG glyph paths, not raster math."""
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_svg import FigureCanvasSVG
    from reportlab.platypus import Flowable
    from reportlab.graphics import renderPDF
    from svglib.svglib import svg2rlg

    if len(latex) > 4000:
        raise DocumentValidationError("Equation exceeds print limits")
    fig = Figure(figsize=(7, 0.75))
    fig.text(.025, .46, "$" + latex.strip().strip("$") + "$", fontsize=16, va="center")
    svg = BytesIO()
    try:
        FigureCanvasSVG(fig).print_svg(svg)
        svg.seek(0)
        drawing = svg2rlg(svg)
        if drawing is None:
            raise ValueError("SVG conversion failed")
    except Exception as exc:
        raise DocumentValidationError("Equation cannot be typeset for PDF: " + str(exc)) from exc

    class EquationFlowable(Flowable):
        def __init__(self):
            super().__init__()
            self.width = min(drawing.width, max_width)
            self.factor = self.width / drawing.width
            self.height = drawing.height * self.factor + 6

        def wrap(self, available_width, available_height):
            return self.width, self.height

        def draw(self):
            self.canv.saveState()
            self.canv.scale(self.factor, self.factor)
            renderPDF.draw(drawing, self.canv, 0, 0)
            self.canv.restoreState()

    return EquationFlowable()


def export_pdf(
    document: dict[str, Any], branding: dict[str, str] | None = None,
    *, audience: str = "student", require_approval: bool = True
) -> bytes:
    """Return A4 PDF; text/equations remain vector-rendered and selectable where supported."""
    if require_approval:
        _approved(document)
    data = validate_document(document, audience=audience)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    )

    brand = _branding(branding)
    font = "Helvetica"
    bold = "Helvetica-Bold"
    fonts_dir = Path("/usr/share/fonts/truetype/dejavu")
    if (fonts_dir / "DejaVuSans.ttf").exists():
        if "SomaDejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("SomaDejaVu", str(fonts_dir / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("SomaDejaVuBold", str(fonts_dir / "DejaVuSans-Bold.ttf")))
        font, bold = "SomaDejaVu", "SomaDejaVuBold"
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("SomNormal", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=17, spaceAfter=9, textColor=colors.HexColor("#162e40"))
    title = ParagraphStyle("SomTitle", parent=normal, fontName=bold, fontSize=23, leading=29, spaceAfter=14)
    heading = ParagraphStyle("SomHeading", parent=normal, fontName=bold, fontSize=15, leading=20, spaceBefore=17, keepWithNext=True)
    sub = ParagraphStyle("SomSub", parent=normal, fontName=bold, fontSize=12, leading=18, spaceBefore=8, keepWithNext=True)
    tiny = ParagraphStyle("SomTiny", parent=normal, fontSize=8, leading=11, textColor=colors.HexColor("#52716d"))
    center = ParagraphStyle("SomCenter", parent=tiny, alignment=TA_CENTER)
    out = BytesIO()
    story: list[Any] = []
    printable = A4[0] - 44 * mm
    story.extend([
        Paragraph(escape(brand["school_name"]), sub),
        Paragraph(escape(data["title"]), title),
        Paragraph("TEACHER EDITION" if audience == "teacher" else "LEARNER EDITION", center),
        Spacer(1, 6 * mm),
    ])
    for section in data["body"]["sections"]:
        if section.get("title"):
            story.append(Paragraph(escape(str(section["title"])), heading))
        for block in section["blocks"]:
            kind = block["type"]
            txt = escape(str(block.get("text", ""))).replace("\n", "<br/>")
            if kind == "page_break":
                story.append(PageBreak())
            elif kind == "heading":
                story.append(Paragraph(txt, sub))
            elif kind == "list":
                for line in str(block.get("text", "")).splitlines():
                    if line.strip():
                        story.append(Paragraph("•  " + escape(line.strip()), normal))
            elif kind == "table":
                cells = [[Paragraph(escape(v), normal) for v in row] for row in block["rows"]]
                table = LongTable(cells, colWidths=[printable / len(cells[0])] * len(cells[0]), repeatRows=1, splitByRow=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("VALIGN", (0,0), (-1,-1), "TOP"), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#cad5da")),
                    ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eaf4f2")),
                    ("LEFTPADDING", (0,0), (-1,-1), 8), ("RIGHTPADDING", (0,0), (-1,-1), 8),
                    ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
                ]))
                story.extend([table, Spacer(1, 9)])
            elif kind == "equation":
                matrix = _matrix_rows(block["latex"])
                if matrix:
                    cells = [[Paragraph(escape(v), normal) for v in row] for row in matrix]
                    table = LongTable(cells, colWidths=[printable / len(cells[0])] * len(cells[0]), hAlign="CENTER")
                    story.append(table)
                else:
                    story.append(_svg_equation(block["latex"], printable))
            elif kind == "reference":
                story.append(Paragraph(txt, tiny))
            elif kind in {"paragraph", "question", "activity", "answer", "marking_guide"}:
                story.append(Paragraph(txt, normal))
            else:
                raise DocumentValidationError("Unrenderable block: " + kind)
            if block.get("citations"):
                story.append(Paragraph("Evidence IDs: " + escape(", ".join(block["citations"])), tiny))
    def footer(canvas, pdf):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#546c75"))
        canvas.drawString(22 * mm, 15 * mm, brand["footer"])
        canvas.drawRightString(A4[0] - 22 * mm, 15 * mm, "Page " + str(pdf.page))
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=22*mm, rightMargin=22*mm, topMargin=23*mm, bottomMargin=23*mm, title=data["title"], author=brand["school_name"])
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
