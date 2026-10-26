"""Bounded multimodal parsers: native PDF, scanned PDF/image, DOCX, HTML.

Outputs are source-referenced and conservative. The engine never uses a generative
model to guess labels, equations, locations, or table cell values.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass
from typing import Any
from uuid import UUID

from services.extraction.models import (
    BlockType, ExtractedAsset, ExtractedDocument, add_block, stable_id,
)
from services.math.equations import analyze_equation, is_equation
from services.visuals.classify import classify_visual, map_metadata

AssetWriter = Callable[[str, bytes], None]
FIGURE_RE = re.compile(r"^\s*(?:fig(?:ure)?\.?|map|diagram|chart|graph)\s*\d*[\w.-]*\s*[:.\-]", re.I)
HEADING_RE = re.compile(r"^(?:chapter|unit|topic|section)\s+\d+\b", re.I)
QUESTION_RE = re.compile(r"^(?:exercise|revision questions|questions|activity)\b", re.I)
BULLET_RE = re.compile(r"^(?:[•\-–]|\d+[.)])\s+")


@dataclass(frozen=True)
class Limits:
    max_bytes: int = 40 * 1024 * 1024
    max_pages: int = 400
    max_assets: int = 250
    max_image_pixels: int = 15_000_000
    ocr_dpi: int = 180


class ExtractionError(ValueError):
    pass


def _box(rect: Any) -> tuple[float, float, float, float] | None:
    if rect is None:
        return None
    return tuple(round(float(n), 2) for n in rect[:4])


def _new_doc(raw: bytes, resource_id: UUID, version_id: UUID) -> ExtractedDocument:
    return ExtractedDocument(resource_id=resource_id, version_id=version_id,
                             sha256=hashlib.sha256(raw).hexdigest())


def _kind(text: str, *, heading: bool = False) -> BlockType:
    if heading or HEADING_RE.match(text):
        return BlockType.heading
    if QUESTION_RE.match(text):
        return BlockType.exercise
    if BULLET_RE.match(text):
        return BlockType.list
    if is_equation(text) and (len(text.split()) <= 24 or text.strip().startswith("$")):
        return BlockType.equation
    return BlockType.paragraph


def _emit_text(doc: ExtractedDocument, text: str, page: int, *,
               bbox=None, confidence=1.0, is_ocr=False, heading=False,
               heading_path=None, metadata=None):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return None
    kind = _kind(text, heading=heading)
    latex = None
    needs_review = is_ocr and confidence < 0.85
    if kind == BlockType.equation:
        analysis = analyze_equation(text, is_ocr=is_ocr)
        latex = analysis.latex
        confidence = min(confidence, analysis.confidence)
        needs_review = needs_review or analysis.needs_review
    return add_block(doc, kind, page, text=text, latex=latex,
                     bbox=bbox, confidence=max(0.0, min(1.0, confidence)),
                     needs_review=needs_review, heading_path=heading_path,
                     metadata=metadata)


def _image_meta(raw: bytes, limit: int) -> tuple[int | None, int | None, int | None]:
    try:
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as image:
            if image.width * image.height > limit:
                return image.width, image.height, None
            dpi = image.info.get("dpi")
            return image.width, image.height, round(float(dpi[0])) if dpi else None
    except Exception:
        return None, None, None


def _asset(doc: ExtractedDocument, raw: bytes, page: int, *,
           caption: str | None, license: str, put_asset: AssetWriter | None,
           bbox=None, limits: Limits = Limits(), metadata: dict | None = None):
    if len(doc.assets) >= limits.max_assets:
        doc.processing_warnings.append("asset_limit_reached")
        doc.partial = True
        return None
    if not raw:
        doc.processing_warnings.append(f"empty_visual_page_{page}")
        return None
    number = len(doc.assets)
    asset_id = stable_id(doc.version_id, "visual", page, number)
    digest = hashlib.sha256(raw).hexdigest()
    suffix = "png" if raw[:8] == b"\x89PNG\r\n\x1a\n" else ("jpg" if raw[:3] == b"\xff\xd8\xff" else "bin")
    key = f"extractions/{doc.version_id}/assets/{asset_id}.{suffix}"
    kind, confidence = classify_visual(caption)
    w, h, dpi = _image_meta(raw, limits.max_image_pixels)
    meta = dict(metadata or {})
    if kind == BlockType.map:
        meta["map"] = map_metadata(caption)
    review = caption is None or confidence < 0.9 or kind == BlockType.map
    asset = ExtractedAsset(id=asset_id, type=kind, page=page,
                           object_key=key, caption=caption,
                           width=w, height=h, dpi=dpi,
                           bbox=bbox, checksum_sha256=digest, license=license,
                           confidence=confidence, needs_review=review, metadata=meta)
    if put_asset:
        put_asset(key, raw)
    doc.assets.append(asset)
    add_block(doc, kind, page, asset_id=asset.id, text=caption,
              bbox=bbox, confidence=confidence, needs_review=review,
              metadata={"object_key": key, "caption_linked": bool(caption)})
    if w and h and (w < 800 or h < 600):
        doc.processing_warnings.append(f"low_resolution_asset_{asset.id}")
    return asset


def _ocr_image(raw: bytes, *, psm=3):
    """Return line text, bounding boxes in pixels and mean Tesseract confidence."""
    try:
        from PIL import Image
        import pytesseract
        from pytesseract import Output
    except ImportError:
        return None
    try:
        with Image.open(io.BytesIO(raw)) as image:
            image.load()
            data = pytesseract.image_to_data(image.convert("RGB"),
                                               config=f"--psm {int(psm)}",
                                               output_type=Output.DICT)
    except Exception:
        return None
    lines: dict[tuple[int, int, int], list[dict]] = defaultdict(list)
    for i, word in enumerate(data.get("text", [])):
        if not word.strip():
            continue
        try:
            conf = float(data["conf"][i]) / 100
        except (ValueError, TypeError):
            continue
        if conf < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines[key].append({
            "text": word, "confidence": min(1.0, conf),
            "bbox": (data["left"][i], data["top"][i],
                     data["left"][i] + data["width"][i],
                     data["top"][i] + data["height"][i]),
        })
    result = []
    for words in lines.values():
        rect = [min(w["bbox"][0] for w in words),
                min(w["bbox"][1] for w in words),
                max(w["bbox"][2] for w in words),
                max(w["bbox"][3] for w in words)]
        result.append((" ".join(w["text"] for w in words),
                       _box(rect), sum(w["confidence"] for w in words) / len(words)))
    return sorted(result, key=lambda x: (x[1][1], x[1][0]))


def _pdf(doc, raw, *, license, put_asset, limits):
    try:
        import fitz
    except ImportError as exc:
        raise ExtractionError("PyMuPDF is required for PDF extraction") from exc
    try:
        pdf = fitz.open(stream=raw, filetype="pdf")
    except Exception as exc:
        raise ExtractionError("Unreadable or corrupt PDF") from exc
    with pdf:
        if pdf.needs_pass:
            raise ExtractionError("Encrypted PDF requires authorization and password handling")
        doc.pages_total = len(pdf)
        if len(pdf) > limits.max_pages:
            doc.partial = True
            doc.processing_warnings.append("pdf_page_limit_reached")
        for page_no in range(min(len(pdf), limits.max_pages)):
            page = pdf[page_no]
            number = page_no + 1
            try:
                entries = page.get_text("dict", sort=True)["blocks"]
                tables = []
                try:
                    for t in page.find_tables().tables:
                        tables.append((_box(t.bbox), [[str(c or "") for c in row]
                                                        for row in t.extract()]))
                except (AttributeError, ValueError):
                    doc.processing_warnings.append(f"table_detection_unavailable_page_{number}")
                events = []
                for b in entries:
                    bbox = _box(b.get("bbox"))
                    if b.get("type") == 1 and b.get("image"):
                        events.append((bbox[1] if bbox else 0, bbox[0] if bbox else 0,
                                       "image", (bytes(b["image"]), bbox)))
                    if b.get("type") != 0:
                        continue
                    lines = b.get("lines", [])
                    text = "\n".join("".join(s.get("text", "") for s in line.get("spans", []))
                                     for line in lines).strip()
                    if not text:
                        continue
                    if any(tbox and bbox and bbox[0] >= tbox[0] and bbox[2] <= tbox[2]
                           and bbox[1] >= tbox[1] and bbox[3] <= tbox[3]
                           for tbox, _ in tables):
                        continue
                    sizes = [s.get("size", 12) for line in lines for s in line.get("spans", [])]
                    heading = bool(sizes and max(sizes) >= 17 and len(text) < 170)
                    events.append((bbox[1] if bbox else 0, bbox[0] if bbox else 0,
                                   "text", (text, bbox, heading)))
                for bbox, cells in tables:
                    events.append((bbox[1] if bbox else 0, bbox[0] if bbox else 0,
                                   "table", (bbox, cells)))
                events.sort(key=lambda x: (x[0], x[1]))
                captions = [(v[0], v[1][0]) for v in events
                            if v[2] == "text" and FIGURE_RE.match(v[3][0])]
                readable = 0
                for top, left, kind, value in events:
                    if kind == "text":
                        text, bbox, heading = value
                        readable += len(text)
                        _emit_text(doc, text, number, bbox=bbox, heading=heading)
                    elif kind == "table":
                        bbox, cells = value
                        add_block(doc, BlockType.table, number, table=cells,
                                  bbox=bbox, confidence=0.9,
                                  needs_review=any(not row for row in cells),
                                  metadata={"source": "pymupdf_table"})
                    elif kind == "image":
                        data, bbox = value
                        nearest = sorted(
                            ((abs(y - (bbox[3] if bbox else top)), c) for y, c in captions),
                            key=lambda x: x[0])
                        caption = nearest[0][1] if nearest and nearest[0][0] < 95 else None
                        _asset(doc, data, number, caption=caption, license=license,
                               put_asset=put_asset, bbox=bbox, limits=limits)
                if readable < 30:
                    pix = page.get_pixmap(dpi=limits.ocr_dpi, alpha=False)
                    if pix.width * pix.height > limits.max_image_pixels:
                        doc.processing_warnings.append(f"ocr_pixel_limit_page_{number}")
                        doc.partial = True
                    else:
                        ocr = _ocr_image(pix.tobytes("png"))
                        if ocr is None:
                            doc.processing_warnings.append(f"ocr_unavailable_page_{number}")
                            doc.partial = True
                        else:
                            scale = 72 / limits.ocr_dpi
                            for text, bbox, conf in ocr:
                                rect = tuple(round(x * scale, 2) for x in bbox)
                                _emit_text(doc, text, number, bbox=rect,
                                           confidence=conf, is_ocr=True)
                doc.pages_processed += 1
            except Exception:
                doc.partial = True
                doc.processing_warnings.append(f"page_processing_failed_{number}")
    return doc


def _docx(doc, raw, *, license, put_asset, limits):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ExtractionError("python-docx is required for DOCX extraction") from exc
    try:
        document = Document(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError("Unreadable or corrupt DOCX") from exc
    doc.pages_total = 1
    doc.pages_processed = 1
    doc.processing_warnings.append("docx_page_numbers_unavailable_without_layout_engine")
    children = list(document.element.body.iterchildren())
    captions = [n.text for n in document.paragraphs if FIGURE_RE.match(n.text)]
    for el in children:
        if el.tag.endswith("}tbl"):
            table = Table(el, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            add_block(doc, BlockType.table, 1, table=rows, confidence=1,
                      metadata={"source": "docx_table", "page_unknown": True})
        elif el.tag.endswith("}p"):
            p = Paragraph(el, document)
            style = p.style.name.lower() if p.style else ""
            _emit_text(doc, p.text, 1, heading=style.startswith(("heading", "title")),
                       metadata={"style": style, "page_unknown": True})
            # OOXML stores math as OMML. Preserve XML exactly; convert only supported expressions.
            from services.math.omml import omml_to_latex
            for math_el in el.xpath(".//m:oMath"):
                original = math_el.xml
                parsed = omml_to_latex(math_el)
                add_block(doc, BlockType.equation, 1, latex=parsed,
                          text=None, confidence=0.8 if parsed else 0.1,
                          needs_review=not bool(parsed),
                          metadata={"omml": original, "page_unknown": True})
            for blip in el.xpath(".//a:blip"):
                embed = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                part = document.part.related_parts.get(embed) if embed else None
                if part and hasattr(part, "blob"):
                    cap = captions[0] if len(captions) == 1 else None
                    _asset(doc, part.blob, 1, caption=cap, license=license,
                           put_asset=put_asset, limits=limits,
                           metadata={"page_unknown": True})
    return doc


def _html(doc, raw, *, license, put_asset, limits):
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise ExtractionError("beautifulsoup4 is required for HTML") from exc
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "iframe", "form", "noscript"]):
        tag.decompose()
    doc.pages_total = doc.pages_processed = 1
    doc.processing_warnings.append("html_has_no_stable_page_coordinates")
    root = soup.body or soup
    for tag in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table",
                              "img", "math"], recursive=True):
        if tag.find_parent(["table", "p", "li", "math"]) is not None:
            continue
        if tag.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in r.find_all(["td", "th"], recursive=False)]
                    for r in tag.find_all("tr")]
            add_block(doc, BlockType.table, 1, table=rows)
        elif tag.name == "img":
            import base64
            src = tag.get("src", "")
            caption = tag.get("alt") or (tag.parent.get_text(" ", strip=True)
                        if tag.parent and tag.parent.name == "figure" else None)
            if src.startswith("data:image/") and ";base64," in src:
                try:
                    image = base64.b64decode(src.split(";base64,", 1)[1], validate=True)
                    if len(image) <= limits.max_bytes:
                        _asset(doc, image, 1, caption=caption, license=license,
                               put_asset=put_asset, limits=limits)
                except (ValueError, Exception):
                    doc.processing_warnings.append("invalid_embedded_html_image")
            else:
                doc.processing_warnings.append("external_html_image_not_fetched")
                add_block(doc, BlockType.figure, 1, text=caption, confidence=0.2,
                          needs_review=True, metadata={"unfetched_src": src[:200]})
        elif tag.name == "math":
            annotation = tag.find("annotation", attrs={"encoding": re.compile("tex", re.I)})
            if annotation:
                eq = analyze_equation(annotation.get_text())
                add_block(doc, BlockType.equation, 1, latex=eq.latex,
                          confidence=eq.confidence, needs_review=eq.needs_review,
                          metadata={"mathml": str(tag)})
            else:
                add_block(doc, BlockType.equation, 1, confidence=0.15, needs_review=True,
                          metadata={"mathml": str(tag)})
        else:
            text = tag.get_text(" ", strip=True)
            if text:
                kind = BlockType.heading if tag.name.startswith("h") else (
                    BlockType.list if tag.name == "li" else _kind(text))
                block = _emit_text(doc, text, 1, heading=kind == BlockType.heading)
                if block and kind == BlockType.list:
                    block.type = BlockType.list
    return doc


def extract_bytes(content: bytes, media_type: str, *,
                  resource_id: UUID, version_id: UUID, license: str = "unknown",
                  put_asset: AssetWriter | None = None, limits: Limits = Limits()) -> ExtractedDocument:
    """Synchronous pure entrypoint called by RQ worker; caller persists JSON/assets."""
    if not isinstance(content, bytes) or not content:
        raise ExtractionError("Empty resource")
    if len(content) > limits.max_bytes:
        raise ExtractionError("Resource exceeds configured extraction limit")
    doc = _new_doc(content, resource_id, version_id)
    mime = media_type.lower().split(";", 1)[0].strip()
    if mime == "application/pdf" or content.startswith(b"%PDF-"):
        return _pdf(doc, content, license=license, put_asset=put_asset, limits=limits)
    if mime in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword"} or (content.startswith(b"PK") and mime.endswith("document")):
        return _docx(doc, content, license=license, put_asset=put_asset, limits=limits)
    if mime in {"text/html", "application/xhtml+xml"}:
        return _html(doc, content, license=license, put_asset=put_asset, limits=limits)
    if mime.startswith("image/"):
        _asset(doc, content, 1, caption=None, license=license,
               put_asset=put_asset, limits=limits)
        ocr = _ocr_image(content)
        doc.pages_total = doc.pages_processed = 1
        if ocr is None:
            doc.partial = True
            doc.processing_warnings.append("image_ocr_unavailable")
        else:
            for text, bbox, conf in ocr:
                _emit_text(doc, text, 1, bbox=bbox, confidence=conf, is_ocr=True)
        return doc
    raise ExtractionError("Unsupported media type")
