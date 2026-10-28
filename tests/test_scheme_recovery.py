import io
from uuid import uuid4
from sqlalchemy import select
from app.core.models import BackgroundJob,extractions
from services.schemes.models import imports
from services.schemes.recovery import text_ready,recover_extractions
from services.extraction.models import ExtractedDocument,ExtractedBlock,SourceRef
from test_master_schemes import seed


def test_visual_limit_preserves_text_but_missing_pages_stay_blocked():
    rid,vid=uuid4(),uuid4()
    doc=ExtractedDocument(resource_id=rid,version_id=vid,sha256='0'*64,pages_total=2,pages_processed=2,partial=True,
        processing_warnings=['asset_limit_reached','low_resolution_asset_123'],blocks=[ExtractedBlock(id=uuid4(),type='paragraph',page=1,order=0,text='P5 Mathematics scheme',confidence=1,source_ref=SourceRef(resource_id=rid,version_id=vid,page=1,order=0))])
    assert text_ready(doc) and doc.partial
    doc.processing_warnings.append('ocr_unavailable_page_2')
    assert not text_ready(doc)
    doc.processing_warnings=['asset_limit_reached'];doc.pages_processed=1
    assert not text_ready(doc)


def test_generic_omml_element_is_preserved():
    from docx import Document
    from lxml import etree
    from services.extraction.engine import extract_bytes
    d=Document();p=d.add_paragraph('Mathematics')
    p._p.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>1+1</m:t></m:r></m:oMath>'))
    raw=io.BytesIO();d.save(raw)
    doc=extract_bytes(raw.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document',resource_id=uuid4(),version_id=uuid4())
    math=next(b for b in doc.blocks if b.type=='equation')
    assert math.latex=='1+1' and 'oMath' in math.metadata['omml']


def test_recovery_retains_progress_and_stops_after_two_attempts(db,accounts,monkeypatch):
    school=accounts[0][0].id;seed(db,school)
    row=db.execute(select(imports).where(imports.c.tenant_id==school)).mappings().one()
    parent=db.get(BackgroundJob,row['job_id']);parent.status='failed';parent.progress=55;parent.error='scheme_extraction_failed'
    child=BackgroundJob(tenant_id=school,kind='extraction',status='failed',progress=20,error='extraction_failed:AttributeError');db.add(child);db.flush()
    db.execute(extractions.insert().values(tenant_id=school,resource_version_id=row['resource_version_id'],job_id=child.id,engine='multimodal-v1'))
    payload={'parts':{'saved':{'answer':'kept'}}}
    db.execute(imports.update().where(imports.c.id==row['id']).values(stage='extraction_failed',payload=payload));db.commit()
    monkeypatch.setattr('services.schemes.recovery.load_document',lambda *args:None)
    assert recover_extractions(db)==1
    assert parent.progress==55 and child.status=='queued'
    saved=db.scalar(select(imports.c.payload).where(imports.c.id==row['id']))
    assert saved['parts']==payload['parts'] and saved['recovery']['attempts']==1
    parent.status=child.status='failed';db.execute(imports.update().where(imports.c.id==row['id']).values(stage='extraction_failed'));db.commit()
    assert recover_extractions(db)==0  # Backoff remains in force.
    saved['recovery']['next_retry_at']='2000-01-01';db.execute(imports.update().where(imports.c.id==row['id']).values(payload=saved));db.commit()
    assert recover_extractions(db)==1
    parent.status=child.status='failed';saved=db.scalar(select(imports.c.payload).where(imports.c.id==row['id']));saved['recovery']['next_retry_at']='2000-01-01'
    db.execute(imports.update().where(imports.c.id==row['id']).values(stage='extraction_failed',payload=saved));db.commit()
    assert recover_extractions(db)==0
